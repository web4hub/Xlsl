from openpyxl import load_workbook
from pathlib import Path

path = Path("/mnt/data/Aura.xlsx")
if not path.exists():
    raise FileNotFoundError(path)

wb = load_workbook(path)
if "Quantum_Optimization" in wb.sheetnames:
    del wb["Quantum_Optimization"]

ws = wb.create_sheet("Quantum_Optimization")

edges = {
    "A-B": 1.0, "A-C": 1.3, "B-D": 1.1, "C-D": 0.9, "D-E": 1.2,
    "E-F": 1.0, "E-G": 1.5, "F-H": 1.2, "G-H": 1.1,
    "C-F": 1.4, "B-G": 1.6
}
paths = [
    "A-B-D-E-F-H", "A-B-D-E-G-H", "A-C-D-E-F-H",
    "A-C-D-E-G-H", "A-C-F-H", "A-B-G-H"
]

def path_cost(p):
    nodes = p.split("-")
    return sum(edges[f"{nodes[i]}-{nodes[i+1]}"] for i in range(len(nodes)-1))

costs = [path_cost(p) for p in paths]
P = 10.0
n = len(paths)

qubo = [[0.0] * n for _ in range(n)]
for i in range(n):
    qubo[i][i] = costs[i] - P
    for j in range(i + 1, n):
        qubo[i][j] = qubo[j][i] = 2 * P

best = min(range(n), key=lambda i: costs[i])
selection = [1 if i == best else 0 for i in range(n)]
objective = (
    sum((costs[i] - P) * selection[i] for i in range(n))
    + sum(qubo[i][j] * selection[i] * selection[j]
          for i in range(n) for j in range(i + 1, n))
)

# Title / metadata
ws.append(["Aura Quantum Optimization"])
ws.append(["Model", "QAOA / QUBO path selection"])
ws.append(["Constraint", "Exactly one path selected"])
ws.append(["QAOA repetitions", 2])
ws.append(["Penalty P", P])
ws.append(["Objective", "Minimize path cost subject to one-hot selection"])
ws.append([])

# Graph
ws.append(["GRAPH", "Edge", "Weight"])
for e, w in edges.items():
    ws.append(["", e, w])
ws.append([])

# Paths
ws.append(["PATHS", "Path", "Cost", "Selected", "QUBO diagonal"])
for i, p in enumerate(paths):
    ws.append([f"x{i}", p, costs[i], selection[i], qubo[i][i]])
ws.append([])

# QUBO
ws.append(["QUBO MATRIX"] + [f"x{i}" for i in range(n)])
for i in range(n):
    ws.append([f"x{i}"] + qubo[i])
ws.append([])

# Result
ws.append(["RESULT", "Value"])
ws.append(["Selected path", paths[best]])
ws.append(["Binary vector", str(selection)])
ws.append(["Path cost", costs[best]])
ws.append(["Objective value", objective])
ws.append([])

# Simulation metadata
ws.append(["SIMULATION METADATA", "Value"])
metadata = [
    ("Problem type", "Predefined path selection / QUBO"),
    ("Variables", n),
    ("QAOA layers (reps)", 2),
    ("Penalty weight", P),
    ("Solver formulation", "QuadraticProgram + MinimumEigenOptimizer + QAOA"),
    ("Execution status", "Model recorded; Qiskit execution requires a Qiskit runtime environment"),
]
for k, v in metadata:
    ws.append([k, v])

# Basic formatting
from openpyxl.styles import Font, PatternFill, Alignment
section_names = {"GRAPH", "PATHS", "QUBO MATRIX", "RESULT", "SIMULATION METADATA"}
for row in ws.iter_rows():
    if row[0].value in section_names or row[0].row == 1:
        for cell in row:
            cell.font = Font(bold=True)
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

for col, width in {"A":24, "B":48, "C":18, "D":14, "E":18, "F":14, "G":14}.items():
    ws.column_dimensions[col].width = width

ws.freeze_panes = "A9"
ws.auto_filter.ref = f"A9:E{ws.max_row}"

wb.save(path)

# Validate
check = load_workbook(path, read_only=True, data_only=True)
qws = check["Quantum_Optimization"]

print("EXECUTION COMPLETE")
print("File:", path)
print("Worksheet:", qws.title)
print("Path costs:", costs)
print("Selected:", paths[best])
print("Binary vector:", selection)
print("Objective:", objective)
print("Workbook sheets:", check.sheetnames)
